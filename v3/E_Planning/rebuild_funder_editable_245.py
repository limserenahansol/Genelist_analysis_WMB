"""Rebuild the editable funder deck around the 245-gene final panel.

Native PowerPoint shapes only (no pictures) so every box stays editable.
Writes the file the user named in Downloads, plus the repo copy.

  1  How the panel was selected
  2  Four rules, with a real IN and OUT
  3  Final 245  -  what it measures
  4  All 245 genes  (identity, activity, TRAP)
  5  All 245 genes  (receptors, morphine state, Dan extras)
  6  BACKUP: 20 populations
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

V3 = Path(__file__).resolve().parents[1]
OUT = V3 / "outputs"
PANEL = OUT / "PANEL_C_expanded_standalone_245genes.xlsx"
BASE = OUT / "xenium_mouse_brain_base_panel.txt"
LONG = OUT / "subclass_markers_expanded" / "Subclass_Discriminating_Markers_long.csv"
PAIRS = OUT / "subclass_markers_all" / "Subclass_Pairwise_Separators.csv"
DL = Path.home() / "Downloads"
DEST_DL = DL / "ORBm_BMAp_panel_funder_EDITABLE.pptx"
DEST_OUT = OUT / "ORBm_BMAp_panel_funder_EDITABLE.pptx"

NAVY, TEAL, RUST, SAGE, GOLD = "1D4E89", "2A6F97", "C44536", "6B7C6A", "B08968"
GREY, CREAM, ICE, WH = "B8B2A8", "F4F1EA", "EAF2F8", "FFFFFF"
INK, MUTED = "1F2937", "5B6472"

GROUPS = [
    ("Cell identity and region",
     ["1_celltype_separator", "3_class_backbone", "11_class_backbone_optional",
      "8_TF_identity", "9_published_regional", "19_expanded_L5IT_separator"], NAVY),
    ("Amygdala extras (Dan)", ["12_GSE283418_added", "18_expanded_BMAp"], GOLD),
    ("Druggable receptor map",
     ["6_GPCR_druggable", "4_GPCR_cell_type_specific",
      "14_Jesse_ORB_GPCR", "17_expanded_ORB_GPCR"], TEAL),
    ("Activity and plasticity", ["5_IEG", "7_plasticity"], SAGE),
    ("Morphine-dependence state",
     ["13_Jesse_morphine_state", "16_expanded_morphine_state"], RUST),
    ("TRAP genetic tag", ["2_reporter_transgene"], GREY),
]


def C(h):
    return RGBColor.from_string(h)


def tbox(s, x, y, w, h, body, size=13, bold=False, colour=INK, align=PP_ALIGN.LEFT,
         space=3, anchor=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    for i, line in enumerate(str(body).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        r.font.color.rgb = C(colour)
    return tb


def rrect(s, x, y, w, h, fill, line, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = C(fill)
    if line:
        sp.line.color.rgb = C(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = True
    return sp


def label_in(sp, text, size=13, bold=False, colour=INK):
    tf = sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        r.font.color.rgb = C(colour)


def connect(s, x1, y1, x2, y2, colour, lw=1.5):
    cn = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    cn.line.color.rgb = C(colour)
    cn.line.width = Pt(lw)


def tile(s, x, y, w, big, lab, col, h=1.10, bigsize=26):
    rrect(s, x, y, w, h, CREAM, col, 1.3)
    tbox(s, x + 0.08, y + 0.04, w - 0.16, 0.52, big, bigsize, True, col, PP_ALIGN.CENTER)
    tbox(s, x + 0.08, y + 0.56, w - 0.16, 0.50, lab, 10.5, False, MUTED, PP_ALIGN.CENTER, space=0)


def bar_chart(s, x, y, w, h, cats, vals, colours, title=None, gap=55):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("genes", vals)
    ch = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd
    ).chart
    ch.has_legend = False
    ch.font.size, ch.font.name = Pt(10), "Calibri"
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(12), True, "Calibri"
        r.font.color.rgb = C(NAVY)
    else:
        ch.has_title = False
    pl = ch.plots[0]
    pl.gap_width = gap
    pl.has_data_labels = True
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(11)
    pl.data_labels.font.bold = True
    pl.data_labels.font.name = "Calibri"
    for i, col in enumerate(colours):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = C(col)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = C(GREY)


def slide(prs, title, sub, footer, tcol=NAVY):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rrect(s, 0, 0, 13.333, 0.12, NAVY, None, shape=MSO_SHAPE.RECTANGLE)
    tbox(s, 0.40, 0.18, 12.5, 0.42, title, 22, True, tcol)
    if sub:
        tbox(s, 0.40, 0.56, 12.5, 0.34, sub, 13, False, MUTED)
    tbox(s, 0.40, 7.16, 12.5, 0.24, footer, 10.5, False, MUTED)
    return s


def pack(genes, n=9):
    genes = list(genes)
    lines = ["    ".join(genes[i:i + n]) for i in range(0, len(genes), n)]
    return "\n".join(lines)


def facts() -> dict:
    sh = pd.read_excel(PANEL, "SHARED_PANEL_ORDER")
    ac = pd.read_excel(PANEL, "ANCHOR_COVERAGE")
    src = pd.read_excel(PANEL, "SOURCES")
    long = pd.read_csv(LONG)

    def bucket(b):
        b = str(b)
        if b.startswith(("13_", "14_", "16_", "17_")):
            return "jesse"
        if b.startswith(("12_", "18_")):
            return "dan"
        return "core"

    sh = sh.copy()
    sh["src"] = sh.block.apply(bucket)
    yld = sh.groupby("src").size()
    lists = {}
    for name, blocks, col in GROUPS:
        lists[name] = list(sh[sh.block.isin(blocks)].gene.astype(str))
    n_listed = sum(len(v) for v in lists.values())
    assert n_listed == len(sh), f"group lists cover {n_listed} of {len(sh)}"
    ac = ac.copy()
    ac["n_sep"] = ac.unique_separators_on_shared_panel.fillna("").apply(
        lambda s: len([x for x in str(s).split(",") if x.strip()])
    )
    return {
        "sh": sh, "ac": ac, "lists": lists, "yld": yld,
        "n": len(sh),
        "n_core": int(yld.get("core", 0)),
        "n_jesse": int(yld.get("jesse", 0)),
        "n_dan": int(yld.get("dan", 0)),
        "n_sources": len(src), "n_doi": int(src.doi.notna().sum()),
        "n_measure": len(long), "n_sub": int(long.subclass.nunique()),
        "n_pairs": len(pd.read_csv(PAIRS)),
        "comp": [(n, len(lists[n]), c) for n, _, c in GROUPS],
    }


def slide_process(prs, F):
    s = slide(
        prs,
        "How the gene panel was selected",
        "Standalone custom Xenium panel  -  no 100-gene add-on cap.  A gene still has to earn its "
        "place: detectable here, and informative for cell type, activity, state, or a receptor map.",
        "ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  1/6",
    )
    tbox(s, 0.42, 0.98, 3.5, 0.26, "WHERE THE CANDIDATES CAME FROM", 10, True, NAVY)
    tbox(s, 4.55, 0.98, 4.3, 0.26, "FOUR DECISION RULES", 10, True, NAVY, PP_ALIGN.CENTER)
    tbox(s, 9.80, 0.98, 3.1, 0.26, "WHAT WE ORDER", 10, True, NAVY, PP_ALIGN.CENTER)

    srcs = [
        ("ALLEN + PUBLISHED MARKERS",
         f"Yao 2023 atlas + {F['n_sources']} documented sources\n"
         "Cell-type markers, IEGs, receptors,\npublished OFC / amygdala genes",
         "re-scored in 226,886 cells", f"{F['n_core']} kept", NAVY),
        ("JESSE NIEHAUS",
         "Morphine-dependence transcriptomics\nin orbitofrontal / prefrontal cortex\n"
         "304 unique DEGs + enriched GPCRs",
         "keep if DE in ≥4 of 12 types", f"{F['n_jesse']} kept", RUST),
        ("BERG & SCHERRER",
         "Published amygdala spatial panel\nGSE283418  ·  98 genes\n"
         "52 already sat in the core",
         "keep if informative in BMAp", f"{F['n_dan']} extras", TEAL),
    ]
    SY, SH_, SG = 1.30, 1.48, 0.16
    for i, (t, b, scr, kep, col) in enumerate(srcs):
        y = SY + i * (SH_ + SG)
        rrect(s, 0.42, y, 3.30, SH_, CREAM, col, 1.6)
        hd = rrect(s, 0.42, y, 3.30, 0.34, col, None, shape=MSO_SHAPE.RECTANGLE)
        label_in(hd, t, 11, True, WH)
        tbox(s, 0.56, y + 0.38, 3.02, 0.78, b, 10, False, INK, space=1)
        tbox(s, 0.56, y + SH_ - 0.30, 2.10, 0.24, scr, 9.5, True, col)
        connect(s, 3.72, y + SH_ / 2, 9.72, 3.45, col, 1.5)
        tbox(s, 3.78, y + SH_ / 2 - 0.14, 1.05, 0.26, kep, 11, True, col)

    rules = [
        ("1", "Does it identify one of the 20 cell types?", NAVY),
        ("2", "Is it a druggable receptor we cannot already see?", TEAL),
        ("3", "Does it report the morphine-dependent state?", RUST),
        ("4", "Will Xenium actually detect it in these regions?", SAGE),
    ]
    RY, RH_, RG = 1.30, 0.88, 0.12
    for i, (n, q, col) in enumerate(rules):
        y = RY + i * (RH_ + RG)
        rrect(s, 4.90, y, 4.55, RH_, WH, col, 1.5)
        ov = rrect(s, 5.02, y + (RH_ - 0.32) / 2, 0.32, 0.32, col, None, shape=MSO_SHAPE.OVAL)
        label_in(ov, n, 12, True, WH)
        tbox(s, 5.44, y + 0.12, 3.90, RH_ - 0.24, q, 12, True, col, space=0,
             anchor=MSO_ANCHOR.MIDDLE)

    tbox(s, 4.90, 5.38, 4.55, 0.22, "Every rule is measured in the Allen atlas", 11, True, NAVY,
         PP_ALIGN.CENTER)
    rrect(s, 4.90, 5.62, 4.55, 0.88, ICE, NAVY, 1.4)
    tbox(s, 5.02, 5.70, 4.31, 0.72,
         f"226,886 cells   ·   {F['n_sub']} subclasses\n"
         f"{F['n_measure']:,} measurements   ·   {F['n_pairs']:,} pairwise tests",
         11, False, INK, PP_ALIGN.CENTER, space=1)

    rrect(s, 9.72, 1.30, 3.20, 5.20, ICE, NAVY, 2.2)
    tbox(s, 9.84, 1.42, 2.96, 0.26, "FINAL DECISION", 11.5, True, NAVY, PP_ALIGN.CENTER)
    tbox(s, 9.84, 1.72, 2.96, 0.72, str(F["n"]), 44, True, NAVY, PP_ALIGN.CENTER)
    tbox(s, 9.84, 2.46, 2.96, 0.24, "genes on one section", 11, False, MUTED, PP_ALIGN.CENTER)
    connect(s, 9.92, 2.82, 12.72, 2.82, GREY, 1.0)
    tbox(s, 9.84, 2.92, 2.96, 1.15,
         f"{F['n_core']} Allen + published markers\n"
         f"{F['n_dan']} Berg & Scherrer extras\n"
         f"{F['n_jesse']} Jesse Niehaus",
         11, False, INK, PP_ALIGN.CENTER, space=2)
    connect(s, 9.92, 4.18, 12.72, 4.18, GREY, 1.0)
    tbox(s, 9.84, 4.28, 2.96, 0.90,
         "standalone custom\nno 100-gene add-on cap\nall 245 are designed probes",
         11, True, NAVY, PP_ALIGN.CENTER, space=1)
    connect(s, 9.92, 5.28, 12.72, 5.28, GREY, 1.0)
    tbox(s, 9.84, 5.38, 2.96, 0.95,
         "2 brain regions\n20 cell populations\n1 section, 1 animal",
         11, True, NAVY, PP_ALIGN.CENTER, space=1)


def slide_rules(prs, F):
    s = slide(
        prs,
        "Why a gene is in  -  and why a similar one is out",
        "Each rule is a measurement with a cutoff.  We did not copy collaborator tables, and we did "
        "not dump the Allen transcriptome.",
        "ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  2/6",
        RUST,
    )
    cards = [
        ("1", "Does it identify the cell type?",
         "Must mark one of our 20 populations more specifically than genes already on the panel",
         "Col23a1   3.7x more specific  -  no equal BMAp marker existed",
         "Nfix   another unique for a type that already has Slc17a7 / Nfib / Rorb", NAVY),
        ("2", "Is it a druggable receptor we cannot see?",
         "ORB-enriched, and it fills a receptor family the panel does not already cover",
         "Chrm1   M1 muscarinic  -  we could only see M2 (Chrm2) before",
         "Adgrd1   enriched in 7 subclasses  -  below the ≥8 cutoff", TEAL),
        ("3", "Does it report the morphine-dependent state?",
         "Differential expression counted only in the 12 ORBm populations we image",
         "Per2   changes in 11 of 12 populations  (broadest gene in the dataset)",
         "Hes1   a TF that changes in only 3 of our 12 populations", RUST),
        ("4", "Will the instrument actually detect it?",
         "Fraction of cells expressing it in Allen ORBm / BMAp; sparse probes give empty maps",
         "Grm8   99% of cells, ≥50% in 8 of 12 populations",
         "Htr3a   Dan leftover at 5% of BMAp cells  -  Xenium would barely see it", SAGE),
    ]
    for i, (n, h, m, k, d, col) in enumerate(cards):
        x, y = 0.42 + (i % 2) * 6.30, 1.02 + (i // 2) * 1.78
        rrect(s, x, y, 6.16, 1.66, WH, col, 1.5)
        ov = rrect(s, x + 0.14, y + 0.12, 0.32, 0.32, col, None, shape=MSO_SHAPE.OVAL)
        label_in(ov, n, 12, True, WH)
        tbox(s, x + 0.54, y + 0.10, 5.48, 0.30, h, 13, True, col, space=0)
        tbox(s, x + 0.16, y + 0.46, 5.86, 0.36, m, 10.5, False, MUTED, space=0)
        tbox(s, x + 0.16, y + 0.88, 5.86, 0.32, "IN      " + k, 10.5, False, SAGE, space=0)
        tbox(s, x + 0.16, y + 1.22, 5.86, 0.32, "OUT    " + d, 10.5, False, RUST, space=0)
    rrect(s, 0.42, 4.70, 12.50, 2.20, ICE, NAVY, 1.2)
    tbox(s, 0.62, 4.82, 12.1, 0.28, "What we did  -  and what we refused to do", 13, True, NAVY)
    tbox(s, 0.62, 5.16, 5.95, 1.58,
         "Allen, published OFC/amygdala markers, Jesse's morphine DEGs, and Dan's 98-gene spatial "
         "panel were all screened.\n\n"
         "Every surviving gene was re-scored in these two regions (226,886 cells). A paper or a "
         "DEG table was never enough on its own.",
         12, False, INK, space=2)
    tbox(s, 6.90, 5.16, 5.95, 1.58,
         f"We are not cutting to a 100-gene vendor cap. The limit is scientific.\n\n"
         f"Jesse: 304 DEGs  →  {F['n_jesse']} kept.  Dan: 98  →  {F['n_dan']} extras.  "
         f"The rest would turn this panel into a transcriptome.",
         12, False, INK, space=2)


def slide_result(prs, F):
    s = slide(
        prs,
        "Final panel  -  245 genes, one section, two regions",
        "Standalone custom panel.  Jesse and Dan included.  All 20 target cell populations remain separable.",
        "ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  3/6",
        TEAL,
    )
    for i, (big, lab, col) in enumerate([
        ("245", "genes on the\nshared panel", NAVY),
        ("20 / 20", "cell populations\nstill separable", TEAL),
        ("no cap", "standalone custom\nall 245 are probes", SAGE),
        ("12 + 8", "populations in\nORBm  +  BMAp", RUST),
    ]):
        tile(s, 0.42 + i * 3.16, 1.00, 2.94, big, lab, col)
    cats = [n for n, _, _ in F["comp"]]
    vals = [v for _, v, _ in F["comp"]]
    cols = [c for _, _, c in F["comp"]]
    bar_chart(s, 0.22, 2.28, 7.15, 4.55, cats, vals, cols, "What the 245 genes do")
    tbox(s, 7.55, 2.32, 5.35, 0.28, "What each cell will tell us", 14, True, NAVY)
    for i, (h, b, col) in enumerate([
        ("Identity", "Which of the 20 populations it is  -  cortical layer, interneuron class, amygdala subtype", NAVY),
        ("Activity", "Whether it was recently active  -  Fos, Arc, and the TRAP2 tag (iCre / tdTomato)", TEAL),
        ("State", "Whether it carries the morphine-dependence signature  -  Per2, Pcsk1, Per1, Camk2g and 50 others", RUST),
        ("Receptors", "Which druggable receptors sit on it  -  opioid, muscarinic, mGlu, Rxfp1 and 42 others", SAGE),
    ]):
        y = 2.72 + i * 0.92
        tbox(s, 7.55, y, 5.35, 0.24, h, 13, True, col, space=0)
        tbox(s, 7.55, y + 0.26, 5.35, 0.60, b, 12, False, INK, space=0)


def gene_card(s, x, y, w, h, title, n, genes, col, per_line=8, gsize=10):
    rrect(s, x, y, w, h, WH, col, 1.4)
    hd = rrect(s, x, y, w, 0.36, col, None, shape=MSO_SHAPE.RECTANGLE)
    label_in(hd, f"{title}     {n}", 13, True, WH)
    tbox(s, x + 0.12, y + 0.44, w - 0.24, h - 0.54, pack(genes, per_line), gsize, False, INK, space=1)


def slide_genes_a(prs, F):
    s = slide(
        prs,
        "The 245 genes  ·  1 of 2",
        "Cell identity, activity / plasticity, and the TRAP tag.  Every name is on the order list.",
        "ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  4/6",
    )
    L = F["lists"]
    gene_card(s, 0.38, 0.96, 12.55, 3.55,
              "Cell identity and region", len(L["Cell identity and region"]),
              L["Cell identity and region"], NAVY, per_line=11, gsize=10)
    gene_card(s, 0.38, 4.62, 8.55, 2.28,
              "Activity and plasticity", len(L["Activity and plasticity"]),
              L["Activity and plasticity"], SAGE, per_line=8, gsize=12)
    gene_card(s, 9.08, 4.62, 3.85, 2.28,
              "TRAP genetic tag", len(L["TRAP genetic tag"]),
              L["TRAP genetic tag"], GREY, per_line=2, gsize=16)


def slide_genes_b(prs, F):
    s = slide(
        prs,
        "The 245 genes  ·  2 of 2",
        "Druggable receptors, morphine-dependence state, and Dan / Scherrer amygdala extras.",
        "ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  5/6",
    )
    L = F["lists"]
    gene_card(s, 0.38, 0.96, 12.55, 2.05,
              "Druggable receptor map", len(L["Druggable receptor map"]),
              L["Druggable receptor map"], TEAL, per_line=10, gsize=11)
    gene_card(s, 0.38, 3.12, 12.55, 2.20,
              "Morphine-dependence state  (Jesse)", len(L["Morphine-dependence state"]),
              L["Morphine-dependence state"], RUST, per_line=10, gsize=11)
    gene_card(s, 0.38, 5.42, 12.55, 1.50,
              "Amygdala extras  (Dan / GSE283418)", len(L["Amygdala extras (Dan)"]),
              L["Amygdala extras (Dan)"], GOLD, per_line=10, gsize=11)


def slide_pops(prs, F):
    s = slide(
        prs,
        "The 20 populations we must tell apart",
        "Target list from the Allen Brain Cell Atlas taxonomy.  All 20 remain separable on the 245-gene panel.",
        "BACKUP  ·  ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  6/6",
    )
    ac = F["ac"]
    for i, (big, lab, col) in enumerate([
        ("12", "ORBm populations\n83,406 Allen cells", NAVY),
        ("8", "BMAp populations\n40,709 Allen cells", TEAL),
        (f"{int(ac.n_cells.sum()):,}", "cells in the 20\ntarget populations", SAGE),
        ("20 / 20", "have a dedicated\nmarker on the panel", RUST),
    ]):
        tile(s, 0.42 + i * 3.16, 1.00, 2.94, big, lab, col)
    for reg, col, x in (("ORBm", NAVY, 0.28), ("BMAp", TEAL, 6.72)):
        g = ac[ac.region == reg]
        bar_chart(
            s, x, 2.28, 6.30, 4.45,
            [r.allen_subclass_anchor.split(" ", 1)[1] for r in g.itertuples()],
            [int(r.n_cells) for r in g.itertuples()],
            [col] * len(g),
            f"{reg}  -  {len(g)} populations, {int(g.n_cells.sum()):,} cells",
            gap=45,
        )


def main() -> None:
    F = facts()
    print(f"245 check: n={F['n']} core={F['n_core']} jesse={F['n_jesse']} dan={F['n_dan']}")
    assert F["n"] == 245
    assert F["n_core"] + F["n_jesse"] + F["n_dan"] == 245

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_process(prs, F)
    slide_rules(prs, F)
    slide_result(prs, F)
    slide_genes_a(prs, F)
    slide_genes_b(prs, F)
    slide_pops(prs, F)

    prs.save(DEST_OUT)
    print("repo", DEST_OUT)
    try:
        prs.save(DEST_DL)
        print("Downloads", DEST_DL)
    except OSError as e:
        alt = DL / "ORBm_BMAp_panel_funder_EDITABLE_245.pptx"
        prs.save(alt)
        print("locked, wrote", alt, type(e).__name__)


if __name__ == "__main__":
    main()
