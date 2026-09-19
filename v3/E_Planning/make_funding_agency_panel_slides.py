"""3 schematic slides for a funding agency: how the panel was chosen, and what it is.

Final decision: standalone custom panel, 245 genes (Jesse + Dan included).
No 100-gene Xenium add-on cap.

1  Process  — define types, score in Allen, build core, add collaborator genes
2  Filter   — Jesse and Dan lists were filtered scientifically, not copied in
3  Result   — 245-gene shared panel, 20/20 types, two regions
"""
from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

V3 = Path(__file__).resolve().parents[1]
OUT = V3 / "outputs"
FIG = OUT / "funding_agency_figures"
FIG.mkdir(exist_ok=True)
DL = Path.home() / "Downloads"

NAVY, TEAL, RUST, GOLD, SAGE, GREY, PALE, INK, MUTED, WHITE = (
    "#1D4E89",
    "#2A6F97",
    "#C44536",
    "#B08968",
    "#6B7C6A",
    "#B8B2A8",
    "#F4F1EA",
    "#1F2937",
    "#5B6472",
    "#FFFFFF",
)


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def set_run(p, text, size=16, bold=False, color=None, name="Calibri", align=None):
    p.clear()
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = name
    r.font.color.rgb = color if color is not None else rgb(INK)
    return r


def add_text(slide, x, y, w, h, text, size=16, bold=False, color=None, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        set_run(p, part, size, bold, color, align=align)
    return tb


def add_lines(slide, x, y, w, h, lines, sizes, bolds=None, colors=None, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    bolds = bolds or [False] * len(lines)
    colors = colors or [rgb(INK)] * len(lines)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        set_run(p, line, sizes[i], bolds[i], colors[i], align=align)
    return tb


def rounded(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    # tighter corners
    try:
        adj = sh.adjustments
        adj[0] = 0.12
    except Exception:
        pass
    return sh


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    return sh


def arrow(slide, x, y, w=0.28, h=0.22, fill=GREY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background()
    return sh


def new_slide(prs, title, sub, n, bar=NAVY):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 0.1, bar)
    add_text(s, 0.4, 0.18, 12.5, 0.42, title, 24, True, rgb(bar))
    add_text(s, 0.4, 0.56, 12.5, 0.32, sub, 13, False, rgb(MUTED))
    add_text(
        s,
        0.4,
        7.18,
        12.5,
        0.24,
        f"ORBm + BMAp custom Xenium panel  |  Allen WMB-10X  |  {n}/3",
        11,
        False,
        rgb(MUTED),
    )
    return s


def fig_composition() -> Path:
    # PANEL_C standalone 245: identity 123, receptors 46, morphine 54, activity 20, TRAP 2
    sizes = [123, 46, 54, 20, 2]
    colors = [NAVY, TEAL, RUST, GOLD, SAGE]
    fig, ax = plt.subplots(figsize=(5.4, 4.4))
    wedges, _ = ax.pie(
        sizes,
        colors=colors,
        startangle=90,
        wedgeprops=dict(width=0.46, edgecolor="white", linewidth=2),
    )
    ax.text(0, 0.12, "245", ha="center", va="center", fontsize=28, color=INK, fontweight="bold", fontfamily="DejaVu Sans")
    ax.text(0, -0.22, "genes", ha="center", va="center", fontsize=12, color="#5B6472", fontfamily="DejaVu Sans")
    ax.legend(
        wedges,
        [
            "Cell identity  123",
            "Receptors  46",
            "Morphine state  54",
            "Activity & plasticity  20",
            "TRAP reporters  2",
        ],
        loc="center left",
        bbox_to_anchor=(0.98, 0.5),
        frameon=False,
        fontsize=9,
        labelspacing=0.55,
    )
    ax.set_xlim(-1.2, 2.6)
    ax.axis("equal")
    fig.tight_layout()
    p = FIG / "F_composition_donut.png"
    fig.savefig(p, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p


def slide_process(prs):
    s = new_slide(
        prs,
        "How the gene panel was chosen",
        "One shared Xenium panel for two regions of the same mouse: orbitofrontal cortex (ORBm) and basomedial amygdala (BMAp).",
        1,
    )
    # constraint strip
    rounded(s, 0.4, 0.95, 12.5, 0.62, PALE, GREY)
    add_lines(
        s,
        0.55,
        1.02,
        12.2,
        0.5,
        [
            "Format    ·    standalone custom Xenium panel     ·     no 100-gene add-on cap     ·     every gene is a designed probe",
            "A gene still has to earn its place: detectable in these two regions, and informative for cell type, activity, state, or receptor map.",
        ],
        [13, 12],
        [True, False],
        [rgb(NAVY), rgb(MUTED)],
    )

    steps = [
        ("1", NAVY, "Map the cell types", "20 populations we must tell apart", "12 in ORBm   +   8 in BMAp\nAllen Brain Cell Atlas taxonomy"),
        ("2", TEAL, "Score in a brain atlas", "226,886 cells, two regions", "Allen WMB-10X\nIs the gene abundant?\nIs it specific to one type?"),
        ("3", SAGE, "Build the core", "Markers + activity + receptors", "Cell-type separators\nIEGs and TRAP tags\nDruggable GPCRs\nAll 20 types separable"),
        ("4", RUST, "Add Jesse + Dan", "Same tests, not whole tables", "Morphine-state genes\nAmygdala spatial genes\nFinal panel: 245 genes"),
    ]
    x0, w, gap = 0.4, 2.72, 0.38
    for i, (num, col, title, sub, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        rounded(s, x, 1.78, w, 4.55, WHITE, col)
        oval = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(1.95), Inches(0.42), Inches(0.42))
        oval.fill.solid()
        oval.fill.fore_color.rgb = rgb(col)
        oval.line.fill.background()
        add_text(s, x + 0.18, 1.98, 0.42, 0.38, num, 16, True, rgb(WHITE), PP_ALIGN.CENTER)
        add_text(s, x + 0.68, 1.98, w - 0.85, 0.4, title, 15, True, rgb(col))
        add_text(s, x + 0.18, 2.48, w - 0.36, 0.55, sub, 13, True, rgb(INK))
        add_text(s, x + 0.18, 3.15, w - 0.36, 2.9, body, 13, False, rgb(MUTED))
        if i < 3:
            arrow(s, x + w + 0.05, 3.85, 0.28, 0.22, GREY)

    add_text(
        s,
        0.4,
        6.5,
        12.5,
        0.55,
        "A gene is not added because it appeared in a paper or a DEG table.  It is added if the atlas says it is detectable here, and it answers a question the rest of the panel cannot.  We do not cut to a 100-gene vendor cap.",
        14,
        False,
        rgb(INK),
    )


def slide_filter(prs):
    s = new_slide(
        prs,
        "Collaborator lists were filtered — not copied in",
        "Two external datasets were tested against the same rule: will the probe be visible, and does it add information we do not already have?",
        2,
        bar=TEAL,
    )

    # Jesse card
    rounded(s, 0.4, 1.0, 6.15, 5.95, WHITE, NAVY)
    rect(s, 0.4, 1.0, 6.15, 0.55, NAVY)
    add_text(s, 0.55, 1.08, 5.85, 0.4, "Jesse   ·   morphine transcriptomics in ORB / prefrontal cortex", 14, True, rgb(WHITE))

    add_lines(
        s,
        0.65,
        1.7,
        5.7,
        0.85,
        ["304 unique DEGs", "→    71 kept"],
        [28, 22],
        [True, True],
        [rgb(NAVY), rgb(RUST)],
    )
    add_text(s, 0.65, 2.65, 5.7, 0.7, "Rule    Keep if it changes in at least 4 of our 12 ORBm cell types, or is an ORB-enriched receptor.  Not a 100-slot cut.", 13, False, rgb(MUTED))

    rounded(s, 0.65, 3.45, 5.7, 1.55, "#EAF2F8", NAVY)
    add_lines(
        s,
        0.8,
        3.52,
        5.4,
        1.4,
        [
            "KEPT",
            "54 morphine-state genes   (Per2, Pcsk1, Per1, Camk2g, …)",
            "17 ORB-enriched receptors   (Chrm1, Grm8, Gpr26, Rxfp1, …)",
            "Includes genes DE in both glutamatergic and GABA types",
        ],
        [12, 13, 13, 13],
        [True, False, False, False],
        [rgb(NAVY), rgb(INK), rgb(INK), rgb(INK)],
    )
    rounded(s, 0.65, 5.15, 5.7, 1.5, PALE, GREY)
    add_lines(
        s,
        0.8,
        5.22,
        5.4,
        1.35,
        [
            "NOT COPIED IN   ·   233 genes",
            "DE in fewer than 4 of our 12 ORBm cell types",
            "Most remaining transcription factors",
            "A panel is not a transcriptome.",
        ],
        [12, 13, 13, 13],
        [True, False, False, False],
        [rgb(RUST), rgb(INK), rgb(INK), rgb(MUTED)],
    )

    # Dan card
    rounded(s, 6.8, 1.0, 6.15, 5.95, WHITE, TEAL)
    rect(s, 6.8, 1.0, 6.15, 0.55, TEAL)
    add_text(s, 6.95, 1.08, 5.85, 0.4, "Dan   ·   amygdala spatial panel   (GSE283418)", 14, True, rgb(WHITE))

    add_lines(
        s,
        7.05,
        1.7,
        5.7,
        0.85,
        ["98 spatial genes", "→    26 extras"],
        [28, 22],
        [True, True],
        [rgb(TEAL), rgb(RUST)],
    )
    add_text(s, 7.05, 2.65, 5.7, 0.7, "Rule    Keep if it is informative in Allen BMAp.  52 of 98 were already in the core, so they are not counted twice.", 13, False, rgb(MUTED))

    rounded(s, 7.05, 3.45, 5.7, 1.55, "#E7F1F4", TEAL)
    add_lines(
        s,
        7.2,
        3.52,
        5.4,
        1.4,
        [
            "KEPT AS EXTRAS",
            "14 stronger BMAp genes  +  12 additional amygdala genes",
            "Includes Sfrp1, Glipr1, Vdr once slots are not scarce",
            "No existing gene was removed or re-ranked",
        ],
        [12, 13, 13, 13],
        [True, False, False, False],
        [rgb(TEAL), rgb(INK), rgb(INK), rgb(INK)],
    )
    rounded(s, 7.05, 5.15, 5.7, 1.5, PALE, GREY)
    add_lines(
        s,
        7.2,
        5.22,
        5.4,
        1.35,
        [
            "NOT COPIED IN   ·   remainder",
            "Not scored in Allen BMAp, or below a low-information cutoff",
            "Genes that only mark distant nuclei",
            "Do not copy the whole 98-gene spatial list.",
        ],
        [12, 13, 13, 13],
        [True, False, False, False],
        [rgb(RUST), rgb(INK), rgb(INK), rgb(MUTED)],
    )


def slide_result(prs, donut: Path):
    s = new_slide(
        prs,
        "Final panel   ·   245 genes, one slide, two regions",
        "Standalone custom Xenium panel for ORBm and BMAp.  Jesse and Dan genes included.  All 20 cell types remain separable.",
        3,
        bar=RUST,
    )

    kpis = [
        (NAVY, "245", "genes on the\nshared panel"),
        (TEAL, "20 / 20", "cell types still\nseparable"),
        (SAGE, "no cap", "standalone custom\nall 245 are probes"),
        (RUST, "2", "regions read from\nthe same mouse"),
    ]
    for i, (col, num, lab) in enumerate(kpis):
        x = 0.4 + i * 3.2
        rounded(s, x, 1.0, 3.0, 1.35, WHITE, col)
        add_text(s, x + 0.12, 1.08, 2.76, 0.55, num, 26, True, rgb(col), PP_ALIGN.CENTER)
        add_text(s, x + 0.12, 1.62, 2.76, 0.62, lab, 12, False, rgb(MUTED), PP_ALIGN.CENTER)

    s.shapes.add_picture(str(donut), Inches(0.25), Inches(2.5), Inches(6.35), Inches(4.35))

    rounded(s, 6.7, 2.5, 6.2, 4.35, WHITE, GREY)
    add_text(s, 6.9, 2.62, 5.85, 0.35, "What each cell can tell us", 16, True, rgb(NAVY))
    rows = [
        ("Identity", "Which of the 20 types it is  (layer, interneuron class, amygdala subtype)"),
        ("Activity", "Whether it was recently active  (Fos / Arc / TRAP2 reporter)"),
        ("State", "Morphine-dependence signature  (Per2, Pcsk1, Per1, Camk2g)"),
        ("Receptors", "Druggable GPCRs on that type  (opioids, muscarinic, mGlu, Rxfp1)"),
    ]
    for i, (head, body) in enumerate(rows):
        y = 3.08 + i * 0.85
        rounded(s, 6.9, y, 5.8, 0.78, PALE)
        add_text(s, 7.05, y + 0.08, 5.5, 0.28, head, 13, True, rgb(RUST))
        add_text(s, 7.05, y + 0.36, 5.5, 0.36, body, 12, False, rgb(INK))


def main() -> None:
    donut = fig_composition()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_process(prs)
    slide_filter(prs)
    slide_result(prs, donut)

    ppt = OUT / "ORBm_BMAp_Xenium_panel_funding_agency.pptx"
    prs.save(ppt)
    dest = DL / "ORBm_BMAp_Xenium_panel_funding_agency.pptx"
    try:
        prs.save(dest)
        print("Downloads", dest)
    except OSError as e:
        alt = DL / "ORBm_BMAp_Xenium_panel_funding_agency_v2.pptx"
        prs.save(alt)
        print("locked, wrote", alt, type(e).__name__)
    print("repo", ppt, ppt.stat().st_size)


if __name__ == "__main__":
    main()
