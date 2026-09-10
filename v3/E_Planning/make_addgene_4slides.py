"""English 4-slide deck: Jesse data (2), Dan data (1), final add list (1).

Uses the same result figures as ORBm_BMAp_Jesse_Dan_panel_decision.pptx
so the slides show how the data look, not only a text summary.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

V3 = Path(__file__).resolve().parents[1]
OUT = V3 / "outputs"
JFIG = OUT / "jesse_orb_figures"
DFIG = OUT / "jesse_dan_figures"
DL = Path.home() / "Downloads"
FONT = "Calibri"

NAVY = RGBColor(0x1D, 0x4E, 0x89)
TEAL = RGBColor(0x2A, 0x6F, 0x97)
RUST = RGBColor(0xC4, 0x45, 0x36)
INK = RGBColor(0x2C, 0x2C, 0x2C)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xF4, 0xF1, 0xEA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_run(p, text, size=16, bold=False, color=INK):
    p.clear()
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return r


def add_bar(slide, color=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def add_footer(slide, n, total=4):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.18), Inches(12.5), Inches(0.26))
    set_run(
        box.text_frame.paragraphs[0],
        f"Xenium ORBm + BMAp  |  same figures as the long collaborator deck  |  {n}/{total}",
        11,
        False,
        MUTED,
    )


def add_title(slide, text, color=NAVY):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.42))
    set_run(box.text_frame.paragraphs[0], text, 24, True, color)


def add_sub(slide, text):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.58), Inches(12.5), Inches(0.32))
    set_run(box.text_frame.paragraphs[0], text, 13, False, MUTED)


def add_table(slide, rows, left, top, width, height, col_w, header=NAVY):
    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)
    ).table
    for i, w in enumerate(col_w):
        table.columns[i].width = Inches(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            set_run(cell.text_frame.paragraphs[0], str(val), 13 if i == 0 else 13, i == 0, WHITE if i == 0 else INK)
            cell.text_frame.word_wrap = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = header if i == 0 else (WHITE if i % 2 == 0 else PALE)


# ----- 1 Jesse IEG data -----
s = prs.slides.add_slide(BLANK)
add_bar(s, NAVY)
add_title(s, "Jesse  |  how the morphine DEGs look")
add_sub(s, "5-day escalating morphine  ·  DESeq2 subclass pseudobulk, padj 0.1  ·  PL–ILA–ORB.  Left: widest IEGs.  Right: which cell types change.")
s.shapes.add_picture(str(JFIG / "Jesse_Fig2_top_IEG_plasticity.png"), Inches(0.25), Inches(0.98), Inches(7.15), Inches(5.95))
s.shapes.add_picture(str(JFIG / "Jesse_Fig3_IEG_by_subclass.png"), Inches(7.45), Inches(0.98), Inches(5.55), Inches(5.95))
add_footer(s, 1)

# ----- 2 Jesse GPCR + IEG matrix -----
s = prs.slides.add_slide(BLANK)
add_bar(s, NAVY)
add_title(s, "Jesse  |  how the ORB GPCRs and IEG matrix look")
add_sub(s, "Left: GPCRs enriched in PL–ILA–ORB vs the rest of the atlas (anatomy, not morphine).  Right: top IEG DEGs × subclass.")
s.shapes.add_picture(str(JFIG / "Jesse_Fig4_enriched_GPCRs.png"), Inches(0.25), Inches(0.98), Inches(6.4), Inches(5.95))
s.shapes.add_picture(str(JFIG / "Jesse_Fig5_IEG_subclass_matrix.png"), Inches(6.75), Inches(0.98), Inches(6.25), Inches(5.95))
add_footer(s, 2)

# ----- 3 Dan data -----
s = prs.slides.add_slide(BLANK)
add_bar(s, TEAL)
add_title(s, "Dan  |  how the 98-gene spatial panel looks", TEAL)
add_sub(s, "GSE283418 Resolve smFISH  ·  98 amygdala genes vs our shared panel, re-scored in Allen BMAp.  Left: fate of 98.  Right: the 14 extras.")
s.shapes.add_picture(str(DFIG / "Dan_Fig1_98gene_fate.png"), Inches(0.25), Inches(0.98), Inches(5.7), Inches(5.95))
s.shapes.add_picture(str(DFIG / "Dan_Fig2_14added_pct.png"), Inches(6.05), Inches(0.98), Inches(6.95), Inches(5.95))
add_footer(s, 3)

# ----- 4 final decision -----
s = prs.slides.add_slide(BLANK)
add_bar(s, RUST)
add_title(s, "Final decision  |  genes added to the panel", RUST)
add_sub(s, "Read from the figures: Jesse’s new IEG program + high-abundance ORB GPCRs. Dan’s 14 are already on MSGS111.")
add_table(
    s,
    [
        ["Decision", "Genes", "Why  (from the data on slides 1–3)", "Slots"],
        [
            "Already on the sheet",
            "Dan 14   Lamb3  Cck  Col23a1 …",
            "38/98 already on the 144. These 14 were the remaining BMAp-informative genes (right plot, slide 3).",
            "14  done",
        ],
        [
            "ADD  free",
            "Rxfp1",
            "Jesse’s top ORB-enriched GPCR (slide 2, left). Already on Xenium Mouse Brain v1. Allen 34% in L5 ET.",
            "0",
        ],
        [
            "ADD",
            "Per2    Pcsk1    Per1",
            "Widest new IEGs on slide 1 — broader than Fos/Arc. Slide 2 matrix: they cut across ORBm excitatory types.",
            "3",
        ],
        [
            "ADD",
            "Chrm1    Grm8",
            "ORB-enriched on slide 2. Allen absolute % is 92% and 99% — unlike Mas1 / Gpr68 / Mchr1 (45 / 40 / 32%).",
            "2",
        ],
        [
            "Not added",
            "Mas1   Gpr68   Mchr1    160 TFs",
            "High on Jesse’s relative-enrichment bar, but too low or too generic for a scarce Xenium slot.",
            "—",
        ],
    ],
    0.35,
    1.02,
    12.6,
    5.15,
    [2.15, 2.85, 6.4, 1.2],
    header=RUST,
)
box = s.shapes.add_textbox(Inches(0.4), Inches(6.28), Inches(12.5), Inches(0.75))
tf = box.text_frame
tf.word_wrap = True
set_run(tf.paragraphs[0], "Final new add:    Rxfp1   +   Per2  Pcsk1  Per1   +   Chrm1  Grm8", 20, True, RUST)
p = tf.add_paragraph()
set_run(
    p,
    "Five custom slots on top of Dan 14  →  119 if every Dan gene is kept.  If the cap stays 100, drop CEA-border Dan genes first.",
    14,
    False,
    INK,
)
add_footer(s, 4)

ppt = OUT / "Add_genes_Jesse_Allen_Dan_4slides.pptx"
prs.save(ppt)
print("saved", ppt.exists(), ppt.stat().st_size)
for dest in (DL / "Add_genes_Dan_Jesse_final_EN.pptx", DL / "Add_genes_Jesse_Allen_Dan_4slides.pptx"):
    try:
        prs.save(dest)
        print("also", dest.name)
    except OSError as e:
        print("skip", dest.name, type(e).__name__)
