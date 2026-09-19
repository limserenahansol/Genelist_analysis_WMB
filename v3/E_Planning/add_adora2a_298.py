"""Add Adora2a -> 298 genes, and update the user's edited deck in place.

WHY Adora2a. It was found by cross-checking the panel against the IUPHAR/DrugBank
drug-target table rather than by the expression screen, and that distinction matters:

  measured here   93.9% of cells in 062 STR D2 Gaba, +40.2pp over every other cell type
                  in the section, and above 50% in only 2 of the 79 types - so it is in
                  the "specific" receptor tier, not a broad one.
  drug            Istradefylline (Nourianz), FDA-approved 2019. Caffeine acts here too.
  why it was not  The genome-wide GPCR sweep filtered to genes strong INSIDE the 20 target
  caught earlier  types. Adora2a's population, 062 STR D2 Gaba, is striatal tissue that
                  falls inside the sAMY section but is not one of the 20 targets, so the
                  filter passed over it. "Zero GPCRs missing" was true for the targets and
                  only for the targets.
  R7             NOT a redundant second marker. 062 STR D2 Gaba currently has no dedicated
                  marker on the panel at all - the best is Drd2 at 95% with only a +12pp
                  margin. Adora2a at +40.2pp becomes its first real one.

HONEST CAVEAT, carried into the workbook: Adora2a is detected in under 20% of cells in
every one of the 20 target cell types. It contributes nothing to ORBm or BMAp identity.
It earns its probe as a drug-target readout on a neighbouring striatal population that
the sAMY section captures, and because A2A-D2 interaction is central to the addiction
literature. If striatal tissue is not of interest, this is the one gene to drop first.

The deck is edited IN PLACE from the user's own 5-slide edit, so their curation survives.
Their footers still read "/7" from the 7-slide version; those are corrected to "/5".
"""
from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from pptx import Presentation

V3 = Path(__file__).resolve().parents[1]
O = V3 / "outputs"
DL = Path(r"C:\Users\hsollim\Downloads")
SRC_X = DL / "FINAL_Xenium_panel_ORBm_BMAp_297genes.xlsx"
DST_X = DL / "FINAL_Xenium_panel_ORBm_BMAp_298genes_FINAL.xlsx"
SRC_P = DL / "ORBm_BMAp_Xenium_panel_297_FINAL_HANSOL.pptx"
DST_P = DL / "ORBm_BMAp_Xenium_panel_298_FINAL_HANSOL.pptx"

GENE = "Adora2a"
BLOCK = "6_GPCR_druggable"
SERVES = "062 STR D2 Gaba :: striatal D2 medium spiny neurons (neighbour population in sAMY)"
WHY = (
    "Adenosine A2A receptor. Target of istradefylline (Nourianz, FDA-approved 2019) and of "
    "caffeine. Detected in 93.9% of cells in 062 STR D2 Gaba, +40.2pp over every other cell "
    "type in the section, and above 50% in only 2 of the 79 types present - so it is a "
    "specific-tier receptor, not a broad one. It also gives 062 STR D2 Gaba its first "
    "dedicated marker: the best the panel carried before was Drd2 at 95% with a margin of "
    "only +12pp. A2A-D2 receptor interaction is a central axis of the addiction literature. "
    "CAVEAT: Adora2a is under 20% in every one of the 20 ORBm/BMAp target cell types, so it "
    "adds nothing to target identity - it is a drug-target readout on striatal tissue that "
    "the sAMY section happens to capture. It was found by cross-checking the IUPHAR drug "
    "table, not by the expression sweep, which had filtered to genes strong inside the 20 "
    "targets."
)


def excel() -> None:
    shutil.copy2(SRC_X, DST_X)
    wb = load_workbook(DST_X)
    for name in ("SHARED_PANEL_ORDER", "BMAp_ORDER"):
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        r = ws.max_row + 1
        vals = {"order_rank": r - 1, "gene": GENE, "block": BLOCK, "serves": SERVES,
                "why": WHY, "max_pct": 93.9, "category": "Druggable receptor map",
                "what_it_tells_you": "Where an FDA-approved drug target sits in this tissue",
                "named_in_the_SOW": "no"}
        for c, h in enumerate(hdr, start=1):
            if h in vals:
                ws.cell(row=r, column=c, value=vals[h])
        print(f"  {name}: added {GENE} at row {r}")

    ws = wb["FOR_MarkGreg"]
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value="ONE GENE ADDED SINCE THE 297 VERSION").font = Font(
        bold=True, color="C44536")
    ws.cell(row=r, column=2, value=(
        "Adora2a, bringing the panel to 298. It is the adenosine A2A receptor, the target of "
        "istradefylline (Nourianz, FDA-approved 2019) and of caffeine. Measured here it is in "
        "93.9% of cells in 062 STR D2 Gaba with a +40.2pp margin over every other cell type, "
        "and above 50% in only 2 of the 79 types - a specific-tier receptor. It also gives "
        "that population its first dedicated marker; the panel's best was previously Drd2 at "
        "a +12pp margin. "
        "Why it was not in the 297: the genome-wide receptor sweep filtered to genes strong "
        "INSIDE the 20 target cell types, and 062 STR D2 Gaba is striatal tissue that the "
        "sAMY section captures but is not one of the 20 targets. It surfaced instead from a "
        "cross-check against the IUPHAR drug-target table. "
        "Honest caveat: Adora2a is under 20% in every one of the 20 target types, so it adds "
        "nothing to ORBm or BMAp cell identity. It is a drug-target readout on neighbouring "
        "striatal tissue. If striatal tissue is not of interest, this is the first gene to "
        "drop."))
    for c in (1, 2):
        ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 120
    wb.save(DST_X)

    md = pd.read_excel(O / "_marker_detail_297.xlsx")
    summ = pd.DataFrame([
        {"class": "GPCR", "screened": 426, "found_in_atlas": 419,
         "expressed_ge50pct_somewhere": 153, "on_the_panel": 50,
         "on_panel_but_empty_map": 0, "missing_and_strong_in_a_TARGET_type": 0,
         "verdict": "COMPLETE for the 20 target cell types: no mouse GPCR reaches 50% in "
                    "one of them with a >=10pp margin and is absent. Separately, a "
                    "cross-check against the IUPHAR drug table added Adora2a, which is "
                    "specific to a NEIGHBOURING striatal population rather than to a "
                    "target - the sweep's target-only filter had passed over it. Eight "
                    "other drug-target genes remain off the panel because they are not "
                    "expressed here (Chrm4 2.7%, Avpr1b 1.7%, Avpr1a 12.3%, Gpr52 13.6%, "
                    "Hcrtr1 26.5%); Tacr3, Ntsr1 and Gpr6 peak outside the two regions."},
        {"class": "Transcription factor", "screened": 1321, "found_in_atlas": 1312,
         "expressed_ge50pct_somewhere": 619, "on_the_panel": 72,
         "on_panel_but_empty_map": 1, "missing_and_strong_in_a_TARGET_type": 2,
         "verdict": "COMPLETE after the no-redundant-marker rule. Bcl6 (+11.8pp, 022 L5 "
                    "ET) and Lhx5 (+10.5pp, 119 Skor1) are weaker second markers for types "
                    "that already carry Npr3 +20.4pp and Skor1 +51.4pp. Npas4 at 45% is "
                    "the activity-gene exemption: the atlas is resting tissue."},
    ])
    with pd.ExcelWriter(DST_X, engine="openpyxl", mode="a",
                        if_sheet_exists="replace") as w:
        summ.to_excel(w, "GPCR_TF_SCREEN", index=False)
        md.to_excel(w, "MARKERS_PER_TYPE", index=False)
    d = pd.read_excel(DST_X, "SHARED_PANEL_ORDER")
    print(f"wrote {DST_X.name}: {len(d)} genes, dupes {int(d.gene.duplicated().sum())}")


def deck() -> None:
    shutil.copy2(SRC_P, DST_P)
    p = Presentation(DST_P)
    n = len(p.slides)
    SWAP = {
        "297": "298",
        "Final panel  -  297 genes, one section, two regions":
            "Final panel  -  298 genes, one section, two regions",
        "What the 297 genes do": "What the 298 genes do",
    }
    edits = 0
    for si, s in enumerate(p.slides, start=1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text.strip()
            # footers: the user's 5-slide edit still carries /7 from the 7-slide version
            if "Allen WMB-10X" in txt and "/" in txt:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if "/" in r.text:
                            r.text = r.text.split("|")[0] + f"|  {si}/{n}" \
                                if "|" in r.text else r.text
                            r.text = r.text.replace("/7", f"/{n}")
                edits += 1
                continue
            if txt in SWAP:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.text.strip() in SWAP:
                            r.text = SWAP[r.text.strip()]
                        elif "297" in r.text:
                            r.text = r.text.replace("297", "298")
                edits += 1
            # GPCR block count 42 -> 43, only where it is the standalone number
            elif txt == "42":
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.text.strip() == "42":
                            r.text = "43"
                edits += 1
            elif "no 298th gene" in txt:
                _set(sh,
                     "Every mouse GPCR (426) and every mouse transcription factor (1,321) "
                     "was scored here, and none is missing from the 20 target cell types. A "
                     "separate cross-check against the FDA drug-target table then added one "
                     "gene, Adora2a, which is specific to neighbouring striatal tissue "
                     "rather than to a target - so the receptor map now also covers the "
                     "drug targets present in this section.")
                edits += 1
            elif "19 specific, 19 intermediate" in txt:
                _set(sh, txt.replace("All 426 mouse GPCRs and 1,321 transcription factors "
                                     "were scored here and none is missing.",
                                     "All 426 mouse GPCRs and 1,321 transcription factors "
                                     "were scored here; none is missing from the 20 target "
                                     "types, and a drug-table cross-check added Adora2a.")
                     .replace("19 specific, 19 intermediate, 11 universal",
                              "20 specific, 19 intermediate, 11 universal"))
                edits += 1
    p.save(DST_P)
    print(f"wrote {DST_P.name}: {n} slides, {edits} shapes edited, "
          f"{sum(1 for s in p.slides for sh in s.shapes if sh.shape_type == 13)} pictures")


def _set(sh, text: str) -> None:
    """Replace a shape's text, keeping the formatting of its first run."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        tf.text = text
        return
    keep = p0.runs[0]
    font = (keep.font.name, keep.font.size, keep.font.bold,
            keep.font.color.rgb if keep.font.color and keep.font.color.type else None)
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    keep.text = text
    keep.font.name, keep.font.size, keep.font.bold = font[0], font[1], font[2]
    if font[3] is not None:
        keep.font.color.rgb = font[3]


if __name__ == "__main__":
    excel()
    deck()
