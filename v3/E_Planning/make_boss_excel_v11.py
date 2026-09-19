"""One-screen, boss-facing workbook for the v11 LEAN 297 panel.

The working file (PANEL_FINAL_v11_ORBm_BMAp_LEAN.xlsx, 24 sheets) is the audit
trail. This is the version you hand to someone who will give it 90 seconds: the
decision first, the gene list second, the reasoning third, and nothing that is not
a decision or a number.

Output: Downloads/PANEL_v11_BOSS_SUMMARY.xlsx
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

V3 = Path(__file__).resolve().parents[1]
O = V3 / "outputs"
DL = Path(r"C:\Users\hsollim\Downloads")
SRC = DL / "PANEL_FINAL_v11_ORBm_BMAp_LEAN.xlsx"
DST = DL / "PANEL_v11_BOSS_SUMMARY.xlsx"

NAVY, TEAL, RUST, CREAM, ICE = "1D4E89", "2A6F97", "C44536", "F4F1EA", "EAF2F8"
INK, MUTED = "1F2937", "5B6472"

TELLS = {
    "1 cell type": "which of the 20 target populations this cell is",
    "2 subtype marker": "which sub-population inside that cell type",
    "3 GPCR": "which druggable receptor sits on that cell type",
    "4 plasticity": "synaptic remodelling state of that cell",
    "5 TF": "the identity programme that defines the cell type",
    "6 TRAP reporter": "was this cell active during volitional opioid seeking",
    "7 IEG": "how recently the cell fired",
    "8 morphine": "does the cell carry the morphine-dependence signature (Jesse DEGs)",
    "9 SOW clock axis": "clock-gene axis the SOW names - exploratory secondary",
}


def main() -> None:
    lean = pd.read_excel(SRC, "FINAL_GENE_LIST")
    cmp_ = pd.read_excel(SRC, "PANEL_COMPARISON")
    exc = pd.read_excel(SRC, "EXCLUDED_measured")
    srcs = pd.read_excel(SRC, "SOURCES")
    pm = pd.read_excel(SRC, "PAPER_MARKER_CHECK")
    rank = pd.read_excel(O / "Jesse_ORB_priority_ranking.xlsx", "ranking_all")
    G = set(lean.gene)

    cats = lean.groupby("category").gene.count()
    n_orbm = len(pd.read_excel(SRC, "ORBm_LIST"))
    n_bmap = len(pd.read_excel(SRC, "BMAp_LIST"))

    # ---------------- 1. START_HERE ----------------
    start = pd.DataFrame({
        "item": [
            "ORDER THIS",
            "Size and format",
            "Regions",
            "What it answers",
            "Cell types callable",
            "Sub-populations callable",
            "SOW-named genes",
            "Detection quality",
            "Evidence",
            "Judgement call 1 - clock genes",
            "Judgement call 2 - non-neuronal probes",
            "Judgement call 3 - panel size",
            "Full audit trail",
        ],
        "detail": [
            f"{len(lean)} genes - sheet PANEL_297 of this file, or FINAL_GENE_LIST in "
            f"PANEL_FINAL_v11_ORBm_BMAp_LEAN.xlsx",
            "Xenium standalone custom panel. The 10x Mouse Brain base panel is NOT included, so every "
            "gene here is a designed probe and nothing is free.",
            f"ORBm (Allen ROI PL-ILA-ORB) and BMAp (Allen ROI sAMY), one section, same mouse. "
            f"{n_orbm} genes serve ORBm, {n_bmap} serve BMAp, the overlap is shared.",
            "For every tdTomato+ (TRAP) cell: which cell type it is, which sub-population inside that "
            "type, which druggable receptors it carries, and whether it carries the morphine-dependence "
            "signature - active vs yoked-passive at the 24 h post-test window.",
            "17 of 20. The other 3 (L4/5 IT, L5 IT, MEA-BST Sox6 GABA) have no single dedicated marker "
            "anywhere in the 32,285-gene atlas and are called from gene combinations - a property of the "
            "transcriptome, not a gap in the panel.",
            "96 of 96 that are markable at all (the previous version reached 66).",
            "8 of 8 - Fos, Arc, Per2, Satb2, Oprm1, Creb1, Adrb1, Htr2a.",
            f"median {lean.allen_pct_own_population.median():.0f} pct of cells detected in the population "
            f"each gene is meant to report; {int((lean.allen_pct_own_population >= 70).sum())} of "
            f"{len(lean)} above 70 pct. Numbers are from 10x scRNA-seq and are an upper bound on Xenium.",
            "Every percentage was computed locally from 226,886 Allen single cells of these two regions "
            "(ORBm 106,122 / BMAp 120,764). No number is quoted from a paper.",
            "10 core clock genes are included because the SOW names \"clock genes like PER2\" as a "
            "readout. The SOW also de-scoped the ZT cohort, so treat this block as exploratory. It is one "
            "block - delete it and the panel is 287.",
            "8 probes (2.7 pct) name astrocytes, oligodendrocytes, OPC, microglia, endothelium, "
            "pericytes, smooth muscle and VLMC. 28 pct of the section is non-neuronal and Xenium segments "
            "every cell, so without them those cells contaminate the neuronal clusters.",
            "297 is the measured plateau, not a budget. Adding the 20 genes we held back (superset, 317) "
            "did not improve cell-type calling; cutting to 237 lost 30 sub-populations.",
            "PANEL_FINAL_v11_ORBm_BMAp_LEAN.xlsx - 24 sheets: per-gene numbers, what was excluded and "
            "why, paper claims checked against the atlas, benchmarks, sources.",
        ],
    })

    # ---------------- 2. PANEL_297 ----------------
    panel = lean[["order_rank", "gene", "category", "allen_pct_own_population",
                  "allen_own_population", "sow_named"]].copy()
    panel["what_it_tells_you"] = panel.category.map(TELLS)
    panel["on_Jesse_morphine_DEG_list"] = panel.gene.isin(set(rank.gene))
    panel = panel.rename(columns={
        "order_rank": "rank", "category": "category (job on the panel)",
        "allen_pct_own_population": "Allen detection % in its own population",
        "allen_own_population": "that population", "sow_named": "named in the SOW"})
    panel = panel[["rank", "gene", "category (job on the panel)", "what_it_tells_you",
                   "Allen detection % in its own population", "that population",
                   "named in the SOW", "on_Jesse_morphine_DEG_list"]]

    # ---------------- 3. HOW_SELECTED ----------------
    rules = pd.DataFrame([
        {"rule": "1. Does it name the cell type?",
         "cutoff (measured in Allen)": ">=50 pct of cells in that population AND >=20pp above the "
                                       "competing populations",
         "IN, with the number": "Vwc2l - 97 pct of L5 NP cells, +71pp over the next type",
         "OUT, with the number": "Cux1 - 72 pct in L2/3 but 94 pct in an amygdala GABA type, so it does "
                                  "not separate",
         "genes kept": int(cats.get("1 cell type", 0))},
        {"rule": "2. Does it name the sub-population inside that type?",
         "cutoff (measured in Allen)": ">=50 pct of the sub-population AND >=20pp above its sibling "
                                       "sub-populations",
         "IN, with the number": "Hpse - 88 pct of Sst_7 (229 cells) vs 8 pct in the rest of Sst",
         "OUT, with the number": "the other 20 candidates of the discovery set - their sub-population "
                                  "already had a marker",
         "genes kept": int(cats.get("2 subtype marker", 0))},
        {"rule": "3. Is it a receptor / TF / plasticity / activity gene we can actually read?",
         "cutoff (measured in Allen)": "detected in >=50 pct of cells somewhere in ORBm or BMAp; broad "
                                       "expression is fine because it is read as a level inside a named "
                                       "cell type",
         "IN, with the number": "Adrb1 - 84 pct, the beta1-adrenergic receptor the SOW names",
         "OUT, with the number": "Sstr4 - 1.4 pct of cells; the probe would return an empty map",
         "genes kept": int(cats.get("3 GPCR", 0) + cats.get("4 plasticity", 0)
                           + cats.get("5 TF", 0) + cats.get("7 IEG", 0))},
        {"rule": "4. Does it report the morphine-dependent state?",
         "cutoff (measured in Allen)": "differentially expressed after 5-day morphine in >=4 of "
                                               "the 12 ORBm cell types AND >=50 pct detected",
         "IN, with the number": "Mbnl2 - DE in 4 ORBm types, 99.8 pct detected",
         "OUT, with the number": "Zbtb40 - DE in 6 types but only 29 pct detected in the types we image",
         "genes kept": int(cats.get("8 morphine", 0))},
        {"rule": "5. Can we tell a neuron from a glial or vascular cell?",
         "cutoff (measured in Allen)": "one marker per non-neuronal class, >=91 pct in its own class",
         "IN, with the number": "Pdgfra - 99.9 pct of OPC, +96pp",
         "OUT, with the number": "second markers for the same class (Ntsr2, Sox10, Csf1r, Col1a1) - "
                                  "redundant",
         "genes kept": 8},
        {"rule": "6. Is it named in the Statement of Work?",
         "cutoff (measured in Allen)": "kept regardless of level - the SOW asks whether it is expressed, "
                                       "so a low answer is still an answer",
         "IN, with the number": "all 8: Fos, Arc, Per2, Satb2, Oprm1, Creb1, Adrb1, Htr2a",
         "OUT, with the number": "sex-determination genes - excluded on your instruction",
         "genes kept": 8},
    ])

    # ---------------- 4. WHAT_WE_EXCLUDED ----------------
    def bucket(w: str) -> str:
        w = str(w)
        if "50 pct detection bar" in w or "empty map" in w:
            return "Detected in under half the cells - the map would be mostly empty"
        if "already has a dedicated marker" in w:
            return "Redundant - a gene already on the panel does the same job"
        if "sex-determination" in w:
            return "Sex-determination genes - excluded on your instruction"
        if "does not hold" in w or "carry L6 CT instead" in w or "other 10 clock genes" in w:
            return "A published marker claim that did not replicate in the Allen cells"
        return "Out of scope or redundant module"

    exc["bucket"] = exc.why_excluded.apply(bucket)
    ex_sum = (exc.groupby("bucket")
              .agg(n_genes=("gene", "size"),
                   examples=("gene", lambda s: ", ".join(sorted(s)[:6]) + (" ..." if len(s) > 6 else "")))
              .reset_index().sort_values("n_genes", ascending=False))
    ex_sum.loc[len(ex_sum)] = ["Published markers tested and failed (Cux1, Scnn1a, Crym, Tcf4, Reln, "
                               "Pax6, Tshz1, Cd36, Fst, Whrn)", len(pm) - int(pm.include.sum()),
                               "0 of 10 passed - see PAPER_MARKER_CHECK in the full file"]

    # ---------------- 5. VERSIONS ----------------
    vs = cmp_[["panel", "genes", "anchors_with_dedicated_marker_of_20",
               "supertypes_with_dedicated_marker_of_96", "SOW_named_genes_of_8",
               "n_under_50pct_own_population", "region_sheet_orphans"]].copy()
    vs.columns = ["version", "genes", "cell types callable /20", "sub-populations callable /96",
                  "SOW genes /8", "probes that would map almost nothing", "documentation errors"]
    vs["verdict"] = ""
    vs.loc[vs.version.str.contains("LEAN 297"), "verdict"] = "ORDER THIS"
    vs.loc[vs.version.str.contains("SUPERSET 317"), "verdict"] = "same capability, 20 extra probes"
    vs.loc[vs.version.str.contains("ULTIMATE"), "verdict"] = "good coverage but misses 2 SOW genes"
    vs.loc[vs.version.str.contains("301"), "verdict"] = "27 near-empty probes, misses 2 SOW genes"
    vs.loc[vs.version.str.contains("C 245|standalone 246|C 246"), "verdict"] = \
        "22 near-empty probes (no detection screen)"

    # ---------------- 6. SOURCES ----------------
    so = srcs[["short_cite", "what_we_used", "doi", "did_we_download_raw_matrix"]].copy()
    so.columns = ["source", "what we used from it", "DOI", "raw data downloaded?"]

    sheets = {"START_HERE": start, "PANEL_297": panel, "HOW_SELECTED": rules,
              "WHAT_WE_EXCLUDED": ex_sum, "VERSIONS": vs, "SOURCES": so}
    with pd.ExcelWriter(DST, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, name, index=False)

    # ---------------- formatting ----------------
    from openpyxl import load_workbook
    wb = load_workbook(DST)
    thin = Side(style="thin", color="D5D9DE")
    for name, ws in ((n, wb[n]) for n in sheets):
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Calibri", size=10, color=INK)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = Border(bottom=thin)
        ws.sheet_view.showGridLines = False

    w1 = wb["START_HERE"]
    w1.column_dimensions["A"].width = 30
    w1.column_dimensions["B"].width = 118
    for r in range(2, w1.max_row + 1):
        w1.cell(r, 1).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    w1.cell(2, 1).fill = PatternFill("solid", fgColor=ICE)
    w1.cell(2, 2).fill = PatternFill("solid", fgColor=ICE)
    w1.cell(2, 2).font = Font(name="Calibri", size=11, bold=True, color=INK)

    w2 = wb["PANEL_297"]
    for col, wd in zip("ABCDEFGH", (7, 12, 20, 46, 16, 30, 12, 14)):
        w2.column_dimensions[col].width = wd
    for r in range(2, w2.max_row + 1):
        v = w2.cell(r, 5).value
        if isinstance(v, (int, float)):
            w2.cell(r, 5).number_format = "0.0"
            w2.cell(r, 5).fill = PatternFill(
                "solid", fgColor="E8F0E6" if v >= 70 else ("FDF3E7" if v >= 50 else "FBE9E7"))
        if w2.cell(r, 7).value is True:
            for c in range(1, 9):
                w2.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=INK)

    for name, widths in (("HOW_SELECTED", (40, 52, 52, 52, 10)),
                         ("WHAT_WE_EXCLUDED", (62, 10, 60)),
                         ("VERSIONS", (34, 8, 13, 16, 10, 16, 14, 34)),
                         ("SOURCES", (26, 86, 30, 22))):
        ws = wb[name]
        for i, wd in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = wd
    ws = wb["VERSIONS"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 8).value == "ORDER THIS":
            for c in range(1, 9):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=ICE)
                ws.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    wb.save(DST)
    print(f"wrote {DST}")
    for n, d in sheets.items():
        print(f"  {n:18s} {d.shape}")


if __name__ == "__main__":
    main()
